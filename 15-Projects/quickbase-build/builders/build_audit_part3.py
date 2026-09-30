import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("/home/ubuntu/parakeet_jotform_audit.xlsx")

NAVY = "2F4E65"; GREEN = "4ABA81"; LIGHT_GRAY = "F2F2F2"; WHITE = "FFFFFF"; GRAY_TAB = "999999"
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
data_font = Font(name="Calibri", size=10)
bold_font = Font(name="Calibri", size=10, bold=True)
green_fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
gray_fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
white_fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
section_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
thin_border = Border(left=Side(style="thin", color="D0D0D0"), right=Side(style="thin", color="D0D0D0"),
                     top=Side(style="thin", color="D0D0D0"), bottom=Side(style="thin", color="D0D0D0"))
wrap = Alignment(wrap_text=True, vertical="top")
FORM_HEADERS = ["Section Name", "Field Name", "Field Type", "Required", "Options/Values",
                "Conditional Logic", "QB Target Field", "FastField?", "Migration Notes"]

def add_form_tab(wb, tab_name, fields):
    ws = wb.create_sheet(title=tab_name[:31])
    ws.sheet_properties.tabColor = GRAY_TAB
    for c, h in enumerate(FORM_HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.freeze_panes = "A2"
    row_idx = 2; data_count = 0
    for f in fields:
        is_section = (f[1] is None)
        for c, val in enumerate(f, 1):
            cell = ws.cell(row=row_idx, column=c, value=val if val is not None else "")
            cell.border = thin_border; cell.alignment = wrap
            if is_section:
                cell.font = bold_font; cell.fill = section_fill
            else:
                cell.font = data_font
                cell.fill = gray_fill if data_count % 2 == 0 else white_fill
                if c == 4 and str(val).strip().lower() == "yes":
                    cell.fill = green_fill
        if not is_section: data_count += 1
        row_idx += 1
    for c in range(1, len(FORM_HEADERS) + 1):
        mx = 12
        for row in ws.iter_rows(min_col=c, max_col=c):
            for cell in row:
                if cell.value: mx = max(mx, min(len(str(cell.value)), 42))
        ws.column_dimensions[get_column_letter(c)].width = mx + 2
    ws.column_dimensions["E"].width = 40; ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 30; ws.column_dimensions["I"].width = 30

# FORM 5: Weekly Maintenance Checklist
def maint_item(section, name, ftype="Checkbox", req="No", opts="Satisfactory / Needs Work", ff="Yes"):
    return (section, name, ftype, req, opts, "None", f"Maintenance > {name}", ff, "Photo + notes per item")

f5 = [
    ("Header", None, None, None, None, None, None, None, None),
    ("Header", "Community", "Dropdown", "Yes", "Community dropdown (different list — OPEN ITEM)", "None", "Maintenance > Community", "Yes", "Reconcile with master community list"),
    ("Header", "Inspection Date", "Date", "Yes", "", "None", "Maintenance > Inspection Date", "Yes", ""),
    ("Header", "Inspector Name", "Text", "Yes", "", "None", "Maintenance > Inspector", "Yes", ""),
    ("Section 1: Grounds/Landscaping", None, None, None, None, None, None, None, None),
    maint_item("Grounds/Landscaping", "Community Walk"),
    maint_item("Grounds/Landscaping", "Litter Hotspots"),
    maint_item("Grounds/Landscaping", "Weeds"),
    maint_item("Grounds/Landscaping", "Tree Trimming"),
    maint_item("Grounds/Landscaping", "Fallen Limbs"),
    maint_item("Grounds/Landscaping", "Full Community Walk Completed"),
    maint_item("Grounds/Landscaping", "Litter Hotspots Identified & Addressed"),
    maint_item("Grounds/Landscaping", "Weeds in Road Cracks Mapped"),
    maint_item("Grounds/Landscaping", "Tree Trimming Needs Identified"),
    maint_item("Grounds/Landscaping", "Fallen Limb Clean-Up Completed/Scheduled"),
    ("Grounds/Landscaping", "General Notes", "Textarea", "No", "", "None", "Maintenance > Grounds Notes", "Yes", "Photo + notes"),
    ("Section 2: Roads/Infrastructure", None, None, None, None, None, None, None, None),
    maint_item("Roads/Infrastructure", "Road Cracks/Potholes Reviewed"),
    maint_item("Roads/Infrastructure", "Drainage/Standing Water Documented"),
    maint_item("Roads/Infrastructure", "Trip Hazards Reviewed"),
    maint_item("Roads/Infrastructure", "Pedestals Spot-Checked"),
    ("Section 3: Structures/Common Areas", None, None, None, None, None, None, None, None),
    maint_item("Structures/Common Areas", "Dumpster Area Deep Clean & Check"),
    maint_item("Structures/Common Areas", "Common Areas Inspected"),
    maint_item("Structures/Common Areas", "Park-Owned Homes Exterior Confirmed"),
    maint_item("Structures/Common Areas", "Vacant Interior Spot-Checks Completed"),
    ("Structures/Common Areas", "General Notes", "Textarea", "No", "", "None", "Maintenance > Structures Notes", "Yes", "Photo + notes"),
    ("Section 4: Admin/Follow-Up", None, None, None, None, None, None, None, None),
    maint_item("Admin/Follow-Up", "Weekly Priorities Sent to CM"),
    maint_item("Admin/Follow-Up", "Vendor Needs Flagged"),
    maint_item("Admin/Follow-Up", "All Items Entered as WOs in RM"),
    ("Admin/Follow-Up", "Additional Comments / Safety Concerns", "Textarea", "No", "", "None", "Maintenance > Safety Notes", "Yes", ""),
]
add_form_tab(wb, "5-Weekly Maint", f5)

# FORM 6: Incident Report
f6 = [
    ("Incident Details", None, None, None, None, None, None, None, None),
    ("Incident Details", "Date of Report", "Date", "No", "", "None", "Incidents > Report Date", "No", ""),
    ("Incident Details", "Date/Time of Incident", "Date", "No", "", "None", "Incidents > Incident DateTime", "No", ""),
    ("Incident Details", "Community/Location", "Dropdown", "No", "65-community dropdown", "None", "Incidents > Community", "No", ""),
    ("Incident Details", "State", "Dropdown", "No", "FL / SC / GA", "None", "Incidents > State", "No", "Could auto-fill from community"),
    ("Incident Details", "Type of Incident", "Text", "No", "", "None", "Incidents > Incident Type", "No", "Consider standardizing as dropdown"),
    ("Incident Details", "Your Name", "Text", "No", "", "None", "Incidents > Reporter Name", "No", ""),
    ("Incident Description", None, None, None, None, None, None, None, None),
    ("Incident Description", "What Happened (Summary)", "Textarea", "No", "", "None", "Incidents > Description", "No", ""),
    ("Incident Description", "How Discovered", "Textarea", "No", "", "None", "Incidents > How Discovered", "No", ""),
    ("Incident Description", "Who Brought to Mgmt Attention", "Text", "No", "", "None", "Incidents > Reported By", "No", ""),
    ("Incident Description", "Did Person Witness First-Hand", "Radio", "No", "Yes / No", "None", "Incidents > First-Hand Witness", "No", ""),
    ("Incident Description", "Who Was Involved/Affected", "Textarea", "No", "", "None", "Incidents > Involved Parties", "No", ""),
    ("Incident Description", "Were There Witnesses", "Radio", "No", "Yes / No", "Shows witness fields if Yes", "Incidents > Has Witnesses", "No", ""),
    ("Incident Description", "Witness 1 Name", "Text", "No", "", "Visible if Witnesses = Yes", "Incidents > Witness 1 Name", "No", ""),
    ("Incident Description", "Witness 1 Contact", "Text", "No", "", "Visible if Witnesses = Yes", "Incidents > Witness 1 Contact", "No", ""),
    ("Incident Description", "Witness 1 Statement", "Textarea", "No", "", "Visible if Witnesses = Yes", "Incidents > Witness 1 Statement", "No", ""),
    ("Incident Description", "Witness 2 Name", "Text", "No", "", "Visible if Witnesses = Yes", "Incidents > Witness 2 Name", "No", ""),
    ("Incident Description", "Witness 2 Contact", "Text", "No", "", "Visible if Witnesses = Yes", "Incidents > Witness 2 Contact", "No", ""),
    ("Incident Description", "Witness 2 Statement", "Textarea", "No", "", "Visible if Witnesses = Yes", "Incidents > Witness 2 Statement", "No", ""),
    ("Incident Response", None, None, None, None, None, None, None, None),
    ("Incident Response", "First Aid Provided", "Radio", "No", "Yes / No / N/A", "None", "Incidents > First Aid", "No", ""),
    ("Incident Response", "Medical Responder Name", "Text", "No", "", "None", "Incidents > Responder", "No", ""),
    ("Incident Response", "Injuries Treated", "Textarea", "No", "", "None", "Incidents > Injuries", "No", ""),
    ("Incident Response", "Treatment Administered", "Textarea", "No", "", "None", "Incidents > Treatment", "No", ""),
    ("Incident Response", "Taken to Hospital", "Radio", "No", "Yes / No", "None", "Incidents > Hospitalized", "No", ""),
    ("Incident Response", "Which Hospital", "Text", "No", "", "Visible if Hospital = Yes", "Incidents > Hospital Name", "No", ""),
    ("Incident Response", "Emergency Services Contacted", "Radio", "No", "Yes / No", "None", "Incidents > Emergency Services", "No", ""),
    ("Incident Response", "Which Services", "Checkbox", "No", "Police / Fire / Ambulance / Other", "Visible if Emergency = Yes", "Incidents > Service Types", "No", ""),
    ("Incident Response", "Law Enforcement Notified", "Radio", "No", "Yes / No", "None", "Incidents > Law Enforcement", "No", ""),
    ("Incident Response", "Officer Names/Badge Numbers", "Text", "No", "", "Visible if Law Enforcement = Yes", "Incidents > Officer Info", "No", ""),
    ("Incident Response", "Case/Report Number", "Text", "No", "", "Visible if Law Enforcement = Yes", "Incidents > Case Number", "No", ""),
    ("Incident Response", "Area Secured/Evacuated", "Radio", "No", "Yes / No / N/A", "None", "Incidents > Area Secured", "No", ""),
    ("Incident Response", "Supervisors Notified", "Radio", "No", "Yes / No", "None", "Incidents > Supervisors Notified", "No", ""),
    ("Incident Response", "Who Notified", "Text", "No", "", "Visible if Supervisors = Yes", "Incidents > Notified Parties", "No", ""),
    ("Incident Response", "Your Role", "Text", "No", "", "None", "Incidents > Reporter Role", "No", ""),
    ("Victim/Involved Party", None, None, None, None, None, None, None, None),
    ("Victim/Involved Party", "Name", "Text", "No", "", "None", "Incidents > Victim Name", "No", ""),
    ("Victim/Involved Party", "Contact", "Text", "No", "", "None", "Incidents > Victim Contact", "No", ""),
    ("Victim/Involved Party", "Role", "Dropdown", "No", "Resident / Employee / Contractor / Visitor / Other", "None", "Incidents > Victim Role", "No", ""),
    ("Victim/Involved Party", "Nature of Involvement", "Textarea", "No", "", "None", "Incidents > Involvement Details", "No", ""),
    ("Victim/Involved Party", "Injuries/Damages", "Textarea", "No", "", "None", "Incidents > Victim Injuries", "No", ""),
    ("Property Damage", None, None, None, None, None, None, None, None),
    ("Property Damage", "Was There Damage", "Radio", "No", "Yes / No", "None", "Incidents > Has Damage", "No", ""),
    ("Property Damage", "Description", "Textarea", "No", "", "Visible if Damage = Yes", "Incidents > Damage Description", "No", ""),
    ("Property Damage", "Owner", "Text", "No", "", "Visible if Damage = Yes", "Incidents > Damage Owner", "No", ""),
    ("Property Damage", "Estimated Cost", "Currency", "No", "", "Visible if Damage = Yes", "Incidents > Damage Cost", "No", ""),
    ("Property Damage", "Photos", "File Upload", "No", "", "Visible if Damage = Yes", "Incidents > Damage Photos", "No", ""),
    ("Investigation & Follow-Up", None, None, None, None, None, None, None, None),
    ("Investigation & Follow-Up", "Investigator Name", "Text", "No", "", "None", "Incidents > Investigator", "No", ""),
    ("Investigation & Follow-Up", "Summary of Findings", "Textarea", "No", "", "None", "Incidents > Findings", "No", ""),
    ("Investigation & Follow-Up", "Contributing Factors / At-Fault", "Textarea", "No", "", "None", "Incidents > Contributing Factors", "No", ""),
    ("Investigation & Follow-Up", "Preventable", "Radio", "No", "Yes / No / Under Investigation", "None", "Incidents > Preventable", "No", ""),
    ("Investigation & Follow-Up", "Prevention Recommendations", "Textarea", "No", "", "None", "Incidents > Prevention Recs", "No", ""),
    ("Investigation & Follow-Up", "Repairs/Corrective Measures", "Textarea", "No", "", "None", "Incidents > Corrective Measures", "No", ""),
    ("Investigation & Follow-Up", "Disciplinary Actions", "Textarea", "No", "", "None", "Incidents > Disciplinary Actions", "No", ""),
    ("Investigation & Follow-Up", "Policy Updates", "Textarea", "No", "", "None", "Incidents > Policy Updates", "No", ""),
    ("Investigation & Follow-Up", "Training Implemented", "Textarea", "No", "", "None", "Incidents > Training", "No", ""),
    ("Investigation & Follow-Up", "Resolution Status", "Dropdown", "No", "Open-Under Investigation / Open-Awaiting Action / In Progress / Resolved / Closed", "None", "Incidents > Status", "No", "Maps to QB pipeline status"),
    ("Supporting Documentation", None, None, None, None, None, None, None, None),
    ("Supporting Documentation", "Photos", "File Upload", "No", "", "None", "Incidents > Photos", "No", ""),
    ("Supporting Documentation", "Video", "File Upload", "No", "", "None", "Incidents > Video", "No", ""),
    ("Supporting Documentation", "Medical Reports", "File Upload", "No", "", "None", "Incidents > Medical Reports", "No", ""),
    ("Supporting Documentation", "Repair Estimates", "File Upload", "No", "", "None", "Incidents > Repair Estimates", "No", ""),
    ("Supporting Documentation", "Communications", "File Upload", "No", "", "None", "Incidents > Communications", "No", ""),
    ("Supporting Documentation", "Additional Notes", "Textarea", "No", "", "None", "Incidents > Additional Notes", "No", ""),
]
add_form_tab(wb, "6-Incident Report", f6)

wb.save("/home/ubuntu/parakeet_jotform_audit.xlsx")
print("Part 3 done — Forms 5-6")
