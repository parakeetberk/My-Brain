import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("/home/ubuntu/parakeet_jotform_audit.xlsx")

NAVY = "2F4E65"
GREEN = "4ABA81"
RED = "F16262"
YELLOW = "FFCC3B"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
GRAY_TAB = "999999"

header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
data_font = Font(name="Calibri", size=10)
bold_font = Font(name="Calibri", size=10, bold=True)
green_fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
gray_fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
white_fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
section_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
thin_border = Border(
    left=Side(style="thin", color="D0D0D0"), right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"), bottom=Side(style="thin", color="D0D0D0"),
)
wrap = Alignment(wrap_text=True, vertical="top")
center_wrap = Alignment(horizontal="center", vertical="top", wrap_text=True)

FORM_HEADERS = ["Section Name", "Field Name", "Field Type", "Required", "Options/Values",
                "Conditional Logic", "QB Target Field", "FastField?", "Migration Notes"]

def add_form_tab(wb, tab_name, fields, is_fastfield_form=False):
    """fields: list of tuples (section, field_name, field_type, required, options, cond_logic, qb_field, fastfield, notes)
       section rows are indicated by field_name being None"""
    ws = wb.create_sheet(title=tab_name[:31])
    ws.sheet_properties.tabColor = GRAY_TAB
    for c, h in enumerate(FORM_HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.freeze_panes = "A2"

    row_idx = 2
    data_row_count = 0
    for f in fields:
        is_section = (f[1] is None)
        for c, val in enumerate(f, 1):
            cell = ws.cell(row=row_idx, column=c, value=val if val is not None else "")
            cell.border = thin_border
            cell.alignment = wrap
            if is_section:
                cell.font = bold_font
                cell.fill = section_fill
            else:
                cell.font = data_font
                cell.fill = gray_fill if data_row_count % 2 == 0 else white_fill
                # Required = Yes → green fill in col 4
                if c == 4 and str(val).strip().lower() == "yes":
                    cell.fill = green_fill
        if not is_section:
            data_row_count += 1
        row_idx += 1

    # Auto-width
    for c in range(1, len(FORM_HEADERS) + 1):
        max_len = 12
        for row in ws.iter_rows(min_col=c, max_col=c):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 42))
        ws.column_dimensions[get_column_letter(c)].width = max_len + 2
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 30
    ws.column_dimensions["I"].width = 30
    return ws

# ============================================================
# FORM 1: Home Listing Confirmation
# ============================================================
f1 = [
    ("—", None, None, None, None, None, None, None, None),  # no sections
    ("(flat form)", "Community", "Dropdown", "No", "65-community dropdown", "None", "Listings > Community", "No", "Map to master community list"),
    ("(flat form)", "Lot/Unit #", "Text", "No", "", "None", "Listings > Lot/Unit #", "No", ""),
    ("(flat form)", "Asset Name", "Text", "No", "", "None", "Listings > Asset Name", "No", ""),
    ("(flat form)", "Listing URL", "Text", "No", "", "None", "Listings > Listing URL", "No", "New field — store published listing link"),
]
# Remove the fake section row
f1_clean = [
    ("(flat form)", "Community", "Dropdown", "No", "65-community dropdown", "None", "Listings > Community", "No", "Map to master community list"),
    ("(flat form)", "Lot/Unit #", "Text", "No", "", "None", "Listings > Lot/Unit #", "No", ""),
    ("(flat form)", "Asset Name", "Text", "No", "", "None", "Listings > Asset Name", "No", ""),
    ("(flat form)", "Listing URL", "Text", "No", "", "None", "Listings > Listing URL", "No", "New field — store published listing link"),
]
add_form_tab(wb, "1-Listing Confirm", f1_clean)

# ============================================================
# FORM 2: Home Listing Request
# ============================================================
f2 = [
    ("Property Information", None, None, None, None, None, None, None, None),
    ("Property Information", "Asset Name", "Text", "Yes", "", "None", "Listings > Asset Name", "No", ""),
    ("Property Information", "Community", "Dropdown", "Yes", "65-community dropdown", "None", "Listings > Community", "No", "Master community list"),
    ("Property Information", "Lot/Unit #", "Text", "Yes", "", "None", "Listings > Lot/Unit #", "No", ""),
    ("Property Information", "Community Type", "Radio", "Yes", "All Age / 55+", "None", "Listings > Community Type", "No", ""),
    ("Home Information", None, None, None, None, None, None, None, None),
    ("Home Information", "Year", "Text", "No", "", "None", "Listings > Year", "No", ""),
    ("Home Information", "Manufacturer", "Text", "No", "", "None", "Listings > Manufacturer", "No", ""),
    ("Home Information", "Home Type", "Radio", "Yes", "MH / RV / Park Model", "None", "Listings > Home Type", "No", ""),
    ("Home Information", "Bedrooms", "Dropdown", "Yes", "Numeric values", "None", "Listings > Bedrooms", "No", ""),
    ("Home Information", "Bathrooms", "Dropdown", "Yes", "Numeric values", "None", "Listings > Bathrooms", "No", ""),
    ("Home Information", "Dimensions", "Text", "No", "", "None", "Listings > Dimensions", "No", ""),
    ("Home Information", "Square Footage", "Number", "No", "", "None", "Listings > Square Footage", "No", ""),
    ("Home Information", "A/C", "Dropdown", "Yes", "Central / Mini Split / Window / None", "None", "Listings > A/C Type", "No", ""),
    ("Home Information", "Laundry Hookups", "Radio", "Yes", "Yes / No", "None", "Listings > Laundry Hookups", "No", ""),
    ("Home Information", "Washer & Dryer", "Radio", "Yes", "Yes / No", "None", "Listings > Washer & Dryer", "No", ""),
    ("Home Information", "Furnished", "Radio", "Yes", "Turnkey / Partial / None", "None", "Listings > Furnished", "No", ""),
    ("Appliances", None, None, None, None, None, None, None, None),
    ("Appliances", "Stove", "Radio", "No", "Yes / No", "None", "Listings > Stove", "No", ""),
    ("Appliances", "Refrigerator", "Radio", "No", "Yes / No", "None", "Listings > Refrigerator", "No", ""),
    ("Appliances", "Microwave", "Radio", "No", "Yes / No", "None", "Listings > Microwave", "No", ""),
    ("Appliances", "Dishwasher", "Radio", "No", "Yes / No", "None", "Listings > Dishwasher", "No", ""),
    ("Signage & Additional Requests", None, None, None, None, None, None, None, None),
    ("Signage & Additional Requests", "Need signs/flags", "Radio", "No", "Yes / No", "None", "Listings > Signage Requested", "No", ""),
    ("Signage & Additional Requests", "Banner qty", "Number", "No", "", "None", "Listings > Banner Qty", "No", ""),
    ("Signage & Additional Requests", "Bandit/Directional qty", "Number", "No", "", "None", "Listings > Bandit Qty", "No", ""),
    ("Signage & Additional Requests", "Flag qty", "Number", "No", "", "None", "Listings > Flag Qty", "No", ""),
    ("Signage & Additional Requests", "Yard Sign qty", "Number", "No", "", "None", "Listings > Yard Sign Qty", "No", ""),
    ("Signage & Additional Requests", "Additional Comments", "Textarea", "No", "", "None", "Listings > Signage Notes", "No", ""),
    ("Listing Details", None, None, None, None, None, None, None, None),
    ("Listing Details", "Listing Date Request", "Date", "Yes", "", "None", "Listings > Listing Date", "No", ""),
    ("Listing Details", "Unit Listing Type", "Radio", "Yes", "MH / RV / Park Model", "None", "Listings > Unit Listing Type", "No", "May duplicate Home Type — confirm"),
    ("Listing Details", "Listing Type", "Radio", "Yes", "For Rent / For Sale", "None", "Listings > Listing Type", "No", "Drives pricing logic"),
    ("Listing Details", "Suggested Monthly Rental Price", "Currency", "Yes", "", "Should hide if Listing Type = For Sale", "Listings > Suggested Rent", "No", "OPEN ITEM: Both price fields required regardless of type"),
    ("Listing Details", "Suggested Sale Price", "Currency", "Yes", "", "Should hide if Listing Type = For Rent", "Listings > Suggested Sale Price", "No", "OPEN ITEM: Both price fields required regardless of type"),
    ("Listing Details", "Home Photos", "File Upload", "No", "", "None", "Listings > Photos (attachment)", "Yes", "FastField for photo capture in field"),
    ("Listing Details", "Additional Highlights", "Textarea", "No", "", "None", "Listings > Highlights", "No", ""),
    ("Requestor Information", None, None, None, None, None, None, None, None),
    ("Requestor Information", "Name", "Text", "Yes", "", "None", "Listings > Requestor Name", "No", ""),
    ("Requestor Information", "Email", "Email", "Yes", "", "None", "Listings > Requestor Email", "No", ""),
    ("Requestor Information", "Phone", "Phone", "Yes", "", "None", "Listings > Requestor Phone", "No", ""),
]
add_form_tab(wb, "2-Listing Request", f2)

# ============================================================
# FORM 3: Offer to Purchase/Rent
# ============================================================
f3 = [
    ("(flat form)", "Community", "Dropdown", "No", "65-community dropdown", "None", "Offers > Community", "No", "Master community list"),
    ("(flat form)", "Lot/Unit #", "Text", "No", "", "None", "Offers > Lot/Unit #", "No", ""),
    ("(flat form)", "Asset Name", "Text", "No", "", "None", "Offers > Asset Name", "No", ""),
    ("(flat form)", "Offer Type", "Radio", "No", "Rent / Buy", "None", "Offers > Offer Type", "No", ""),
    ("(flat form)", "Initial Listing Date", "Date", "No", "", "None", "Offers > Initial Listing Date", "No", ""),
    ("(flat form)", "Original Price", "Currency", "No", "", "None", "Offers > Original Price", "No", ""),
    ("(flat form)", "Current Asking Price", "Currency", "No", "", "None", "Offers > Current Asking Price", "No", ""),
    ("(flat form)", "Amount Offered", "Currency", "No", "", "None", "Offers > Amount Offered", "No", ""),
    ("(flat form)", "Terms of Offer", "Textarea", "No", "", "None", "Offers > Terms", "No", ""),
    ("(flat form)", "Name of Purchaser", "Text", "No", "", "None", "Offers > Purchaser Name", "No", ""),
    ("(flat form)", "Upload Signed PSA", "File Upload", "No", "", "None", "Offers > PSA Attachment", "No", "OPEN ITEM: QB attachment or SharePoint?"),
    ("(flat form)", "Application Submitted", "Radio", "No", "Yes / No", "None", "Offers > App Submitted", "No", ""),
    ("(flat form)", "Agent Name", "Text", "No", "", "None", "Offers > Agent Name", "No", ""),
    ("(flat form)", "Agent Email", "Email", "No", "", "None", "Offers > Agent Email", "No", ""),
    ("(flat form)", "Regional Manager Email", "Email", "No", "", "None", "Offers > RM Email", "No", "Auto-populate from community relationship"),
]
add_form_tab(wb, "3-Offer Purchase", f3)

# ============================================================
# FORM 4: Utility Request
# ============================================================
f4 = [
    ("(flat form)", "Community", "Dropdown", "No", "65-community dropdown", "None", "Utility Requests > Community", "No", "Master community list"),
    ("(flat form)", "Lot/Unit #", "Text", "No", "", "None", "Utility Requests > Lot/Unit #", "No", ""),
    ("(flat form)", "Street Address", "Text", "No", "", "None", "Utility Requests > Street", "No", ""),
    ("(flat form)", "City", "Text", "No", "", "None", "Utility Requests > City", "No", ""),
    ("(flat form)", "State", "Text", "No", "", "None", "Utility Requests > State", "No", ""),
    ("(flat form)", "Zip", "Text", "No", "", "None", "Utility Requests > Zip", "No", ""),
    ("(flat form)", "Turn On / Off", "Dropdown", "No", "Turn On / Turn Off", "None", "Utility Requests > Action", "No", ""),
    ("(flat form)", "Utility Type", "Checkbox", "No", "Water / Electric / Sewer / Internet", "None", "Utility Requests > Utility Type", "No", "Multi-select"),
    ("(flat form)", "Account # / Meter #", "Text", "No", "", "None", "Utility Requests > Account #", "No", ""),
    ("(flat form)", "Vendor", "Text", "No", "", "None", "Utility Requests > Vendor", "No", ""),
    ("(flat form)", "Reason for Request", "Textarea", "No", "", "None", "Utility Requests > Reason", "No", ""),
    ("(flat form)", "Date Needed", "Date", "No", "", "None", "Utility Requests > Date Needed", "No", ""),
    ("(flat form)", "Requestor Name", "Text", "No", "", "None", "Utility Requests > Requestor", "No", ""),
    ("(flat form)", "Email", "Email", "No", "", "None", "Utility Requests > Email", "No", ""),
    ("(flat form)", "Phone", "Phone", "No", "", "None", "Utility Requests > Phone", "No", ""),
]
add_form_tab(wb, "4-Utility Request", f4)

wb.save("/home/ubuntu/parakeet_jotform_audit.xlsx")
print("Part 2 done — Forms 1-4")
