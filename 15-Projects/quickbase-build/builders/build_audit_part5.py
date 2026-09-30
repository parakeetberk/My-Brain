import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("/home/ubuntu/parakeet_jotform_audit.xlsx")

NAVY = "2F4E65"; GREEN = "4ABA81"; LIGHT_GRAY = "F2F2F2"; WHITE = "FFFFFF"; GRAY_TAB = "999999"
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
data_font = Font(name="Calibri", size=10)
bold_font = Font(name="Calibri", size=10, bold=True)
green_fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
gray_fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
white_fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
section_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
thin_border = Border(left=Side(style="thin", color="D0D0D0"), right=Side(style="thin", color="D0D0D0"),
                     top=Side(style="thin", color="D0D0D0"), bottom=Side(style="thin", color="D0D0D0"))
wrap = Alignment(wrap_text=True, vertical="top")
FORM_HEADERS = ["Section Name", "Field Name", "Field Type", "Required", "Options/Values",
                "Conditional Logic", "QB Target Field", "FastField?", "Migration Notes"]

def add_form_tab(wb, tab_name, fields):
    ws = wb.create_sheet(title=tab_name[:31])
    ws.sheet_properties.tabColor = GRAY_TAB
    for c, h in enumerate(FORM_HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.freeze_panes = "A2"
    row_idx = 2; dc = 0
    for f in fields:
        is_sec = (f[1] is None)
        for c, val in enumerate(f, 1):
            cell = ws.cell(row=row_idx, column=c, value=val if val is not None else "")
            cell.border = thin_border; cell.alignment = wrap
            if is_sec: cell.font = bold_font; cell.fill = section_fill
            else:
                cell.font = data_font; cell.fill = gray_fill if dc % 2 == 0 else white_fill
                if c == 4 and str(val).strip().lower() == "yes": cell.fill = green_fill
        if not is_sec: dc += 1
        row_idx += 1
    for c in range(1, 10):
        mx = 12
        for row in ws.iter_rows(min_col=c, max_col=c):
            for cell in row:
                if cell.value: mx = max(mx, min(len(str(cell.value)), 42))
        ws.column_dimensions[get_column_letter(c)].width = mx + 2
    ws.column_dimensions["E"].width = 40; ws.column_dimensions["F"].width = 35
    ws.column_dimensions["G"].width = 30; ws.column_dimensions["I"].width = 30

# FORM 9: Move Out Inspection — rating-based inspection
def insp(sec, name, req="No"):
    return (sec, name, "Dropdown", req, "Excellent / Good / Fair / Poor / N/A", "None", f"Inspections > {name}", "Yes", "Rating + notes per item")

f9 = [
    ("Header", None, None, None, None, None, None, None, None),
    ("Header", "Community", "Dropdown", "No", "65-community dropdown", "None", "Inspections > Community", "Yes", ""),
    ("Header", "Lot/Unit #", "Text", "No", "", "None", "Inspections > Lot/Unit #", "Yes", ""),
    ("Header", "Date of Inspection", "Date", "No", "", "None", "Inspections > Date", "Yes", ""),
    ("Header", "Inspection Conducted By", "Text", "No", "", "None", "Inspections > Inspector", "Yes", ""),
    ("Exterior", None, None, None, None, None, None, None, None),
    insp("Exterior", "Driveway"), insp("Exterior", "Walkway"), insp("Exterior", "Siding"),
    insp("Exterior", "Skirting"), insp("Exterior", "Utility Connections"),
    insp("Exterior", "Landscaping"), insp("Exterior", "Trees & Shrubs"),
    ("Main Entrance", None, None, None, None, None, None, None, None),
    insp("Main Entrance", "Front Steps/Landing"), insp("Main Entrance", "Front Door"),
    insp("Main Entrance", "Door Hardware"), insp("Main Entrance", "Entrance Threshold"),
    ("Living Room", None, None, None, None, None, None, None, None),
    insp("Living Room", "Flooring"), insp("Living Room", "Walls"), insp("Living Room", "Ceiling"),
    insp("Living Room", "Windows"), insp("Living Room", "Lighting Fixtures"),
    insp("Living Room", "Electrical Outlets"),
    ("Dining Room", None, None, None, None, None, None, None, None),
    insp("Dining Room", "Flooring"), insp("Dining Room", "Walls"), insp("Dining Room", "Ceiling"),
    insp("Dining Room", "Windows"), insp("Dining Room", "Lighting"), insp("Dining Room", "Electrical Outlets"),
    ("Kitchen", None, None, None, None, None, None, None, None),
    insp("Kitchen", "Range"), insp("Kitchen", "Refrigerator"), insp("Kitchen", "Microwave"),
    insp("Kitchen", "Sink"), insp("Kitchen", "Faucet"), insp("Kitchen", "Flooring"),
    insp("Kitchen", "Walls"), insp("Kitchen", "Ceiling"), insp("Kitchen", "Windows"),
    insp("Kitchen", "Lighting"), insp("Kitchen", "Electrical Outlets"), insp("Kitchen", "Cabinets"),
    insp("Kitchen", "Exhaust Fan"), insp("Kitchen", "Fire Alarm/Smoke Detector"),
    ("Bedroom 1", None, None, None, None, None, None, None, None),
    ("Bedroom 1", "Door/Hardware", "Dropdown", "No", "Excellent / Good / Fair / Poor / N/A", "Conditional — based on # of bedrooms", "Inspections > BR1 Door", "Yes", "Bedrooms 2-4 same pattern"),
    insp("Bedroom 1", "Flooring"), insp("Bedroom 1", "Walls"), insp("Bedroom 1", "Ceiling"),
    insp("Bedroom 1", "Windows"), insp("Bedroom 1", "Window Treatments"),
    insp("Bedroom 1", "Closet"), insp("Bedroom 1", "Lighting"), insp("Bedroom 1", "Electrical Outlets"),
    ("Bedrooms 2-4", None, None, None, None, None, None, None, None),
    ("Bedrooms 2-4", "(Same fields as Bedroom 1 × 3)", "Dropdown", "No", "Excellent / Good / Fair / Poor / N/A", "Conditional — each bedroom shown based on count", "Inspections > BR2-4 fields", "Yes", "Repeating section in FastField"),
    ("Bathroom 1", None, None, None, None, None, None, None, None),
    insp("Bathroom 1", "Sink"), insp("Bathroom 1", "Faucet"), insp("Bathroom 1", "Toilet"),
    insp("Bathroom 1", "Shower/Tub"), insp("Bathroom 1", "Towel Rack"),
    insp("Bathroom 1", "Door/Hardware"), insp("Bathroom 1", "Floor"), insp("Bathroom 1", "Walls"),
    insp("Bathroom 1", "Ceiling"), insp("Bathroom 1", "Windows"), insp("Bathroom 1", "Window Treatments"),
    insp("Bathroom 1", "Closet"), insp("Bathroom 1", "Cabinets"), insp("Bathroom 1", "Exhaust Fan"),
    insp("Bathroom 1", "Lighting"), insp("Bathroom 1", "Electrical Outlets"),
    ("Bathrooms 2-4", None, None, None, None, None, None, None, None),
    ("Bathrooms 2-4", "(Same fields as Bathroom 1 × 3)", "Dropdown", "No", "Excellent / Good / Fair / Poor / N/A", "Conditional — each bathroom shown based on count", "Inspections > BA2-4 fields", "Yes", "Repeating section in FastField"),
    ("Other Equipment", None, None, None, None, None, None, None, None),
    insp("Other Equipment", "Heating Equipment"), insp("Other Equipment", "A/C Units"),
    insp("Other Equipment", "Hot Water Heater"), insp("Other Equipment", "Smoke/Fire Alarms"),
    insp("Other Equipment", "Thermostat"), insp("Other Equipment", "Doorbell"),
    ("Sign-Off", None, None, None, None, None, None, None, None),
    ("Sign-Off", "Comments / Notes", "Textarea", "No", "", "None", "Inspections > Notes", "Yes", ""),
    ("Sign-Off", "Inspector Signature", "Signature", "No", "", "None", "Inspections > Signature", "Yes", ""),
]
add_form_tab(wb, "9-Move Out Insp", f9)

# FORM 10: Resident Information Sheet
f10 = [
    ("(flat form)", "Community", "Dropdown", "No", "65-community dropdown", "None", "Residents > Community", "No", ""),
    ("(flat form)", "Lot/Unit #", "Text", "No", "", "None", "Residents > Lot/Unit #", "No", ""),
    ("Resident 1", None, None, None, None, None, None, None, None),
    ("Resident 1", "Name", "Text", "No", "", "None", "Residents > Resident 1 Name", "No", "Sync to Rent Manager"),
    ("Resident 1", "Phone", "Phone", "No", "", "None", "Residents > Resident 1 Phone", "No", ""),
    ("Resident 1", "Email", "Email", "No", "", "None", "Residents > Resident 1 Email", "No", ""),
    ("Resident 2", None, None, None, None, None, None, None, None),
    ("Resident 2", "Name", "Text", "No", "", "None", "Residents > Resident 2 Name", "No", ""),
    ("Resident 2", "Phone", "Phone", "No", "", "None", "Residents > Resident 2 Phone", "No", ""),
    ("Resident 2", "Email", "Email", "No", "", "None", "Residents > Resident 2 Email", "No", ""),
    ("Resident 3", None, None, None, None, None, None, None, None),
    ("Resident 3", "Name", "Text", "No", "", "None", "Residents > Resident 3 Name", "No", ""),
    ("Resident 3", "Phone", "Phone", "No", "", "None", "Residents > Resident 3 Phone", "No", ""),
    ("Resident 3", "Email", "Email", "No", "", "None", "Residents > Resident 3 Email", "No", ""),
    ("Resident 4", None, None, None, None, None, None, None, None),
    ("Resident 4", "Name", "Text", "No", "", "None", "Residents > Resident 4 Name", "No", ""),
    ("Resident 4", "Phone", "Phone", "No", "", "None", "Residents > Resident 4 Phone", "No", ""),
    ("Resident 4", "Email", "Email", "No", "", "None", "Residents > Resident 4 Email", "No", ""),
    ("Vehicle 1", None, None, None, None, None, None, None, None),
    ("Vehicle 1", "Make", "Text", "No", "", "None", "Residents > Vehicle 1 Make", "No", ""),
    ("Vehicle 1", "Model", "Text", "No", "", "None", "Residents > Vehicle 1 Model", "No", ""),
    ("Vehicle 1", "Plate #", "Text", "No", "", "None", "Residents > Vehicle 1 Plate", "No", ""),
    ("Vehicle 2", None, None, None, None, None, None, None, None),
    ("Vehicle 2", "Make", "Text", "No", "", "None", "Residents > Vehicle 2 Make", "No", ""),
    ("Vehicle 2", "Model", "Text", "No", "", "None", "Residents > Vehicle 2 Model", "No", ""),
    ("Vehicle 2", "Plate #", "Text", "No", "", "None", "Residents > Vehicle 2 Plate", "No", ""),
    ("Home Info", None, None, None, None, None, None, None, None),
    ("Home Info", "Year", "Text", "No", "", "None", "Residents > Home Year", "No", ""),
    ("Home Info", "Manufacturer", "Text", "No", "", "None", "Residents > Home Manufacturer", "No", ""),
    ("Home Info", "VIN #", "Text", "No", "", "None", "Residents > VIN", "No", ""),
    ("Home Info", "Copy of Title", "File Upload", "No", "", "None", "Residents > Title Upload", "No", ""),
    ("Home Info", "Insurance Company", "Text", "No", "", "None", "Residents > Insurance Co", "No", ""),
    ("Home Info", "Policy #", "Text", "No", "", "None", "Residents > Policy #", "No", ""),
    ("Contact Info", None, None, None, None, None, None, None, None),
    ("Contact Info", "Name", "Text", "No", "", "None", "Residents > Contact Name", "No", ""),
    ("Contact Info", "Email", "Email", "No", "", "None", "Residents > Contact Email", "No", ""),
    ("Contact Info", "Phone", "Phone", "No", "", "None", "Residents > Contact Phone", "No", ""),
]
add_form_tab(wb, "10-Resident Info", f10)

wb.save("/home/ubuntu/parakeet_jotform_audit.xlsx")
print("Part 5 done — Forms 9-10")
