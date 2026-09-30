import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("/home/ubuntu/parakeet_jotform_audit.xlsx")

NAVY = "2F4E65"; GREEN = "4ABA81"; LIGHT_GRAY = "F2F2F2"; WHITE = "FFFFFF"; GRAY_TAB = "999999"
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
data_font = Font(name="Calibri", size=10)
bold_font = Font(name="Calibri", size=10, bold=True)
green_fill = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
gray_fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
white_fill = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
section_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
thin_border = Border(left=Side(style="thin", color="D0D0D0"), right=Side(style="thin", color="D0D0D0"),
                     top=Side(style="thin", color="D0D0D0"), bottom=Side(style="thin", color="D0D0D0"))
wrap = Alignment(wrap_text=True, vertical="top")
FORM_HEADERS = ["Section Name", "Field Name", "Field Type", "Required", "Options/Values",
                "Conditional Logic", "QB Target Field", "FastField?", "Migration Notes"]

def add_form_tab(wb, tab_name, fields):
    ws = wb.create_sheet(title=tab_name[:31])
    ws.sheet_properties.tabColor = GRAY_TAB
    for c, h in enumerate(FORM_HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.freeze_panes = "A2"
    row_idx = 2; dc = 0
    for f in fields:
        is_sec = (f[1] is None)
        for c, val in enumerate(f, 1):
            cell = ws.cell(row=row_idx, column=c, value=val if val is not None else "")
            cell.border = thin_border; cell.alignment = wrap
            if is_sec: cell.font = bold_font; cell.fill = section_fill
            else:
                cell.font = data_font; cell.fill = gray_fill if dc % 2 == 0 else white_fill
                if c == 4 and str(val).strip().lower() == "yes": cell.fill = green_fill
        if not is_sec: dc += 1
        row_idx += 1
    for c in range(1, 10):
        mx = 12
        for row in ws.iter_rows(min_col=c, max_col=c):
            for cell in row:
                if cell.value: mx = max(mx, min(len(str(cell.value)), 42))
        ws.column_dimensions[get_column_letter(c)].width = mx + 2
    ws.column_dimensions["E"].width = 40; ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 30; ws.column_dimensions["I"].width = 30

# FORM 7: Monthly Equipment & Inventory
def equip(sec, name, ftype="Checkbox", req="No", opts="Yes / No", notes=""):
    return (sec, name, ftype, req, opts, "None", f"Equipment > {name}", "Yes", notes)

f7 = [
    ("Property & Technician Info", None, None, None, None, None, None, None, None),
    ("Property & Technician Info", "Community/Property", "Dropdown", "No", "Community 1/2/3/4/Other — PLACEHOLDER", "None", "Equipment > Community", "Yes", "OPEN ITEM: Needs real community list"),
    ("Property & Technician Info", "Technician Name", "Text", "No", "", "None", "Equipment > Technician", "Yes", ""),
    ("Property & Technician Info", "Inspection Date", "Date", "No", "", "None", "Equipment > Inspection Date", "Yes", ""),
    ("Property & Technician Info", "Month Being Inspected", "Dropdown", "No", "Jan–Dec", "None", "Equipment > Month", "Yes", ""),
    ("Landscaping Equipment", None, None, None, None, None, None, None, None),
    equip("Landscaping Equipment", "Mower — Blades Sharp"),
    equip("Landscaping Equipment", "Mower — Oil Level"),
    equip("Landscaping Equipment", "Mower — Cleaned"),
    equip("Landscaping Equipment", "Mower — Greased"),
    equip("Landscaping Equipment", "Mower — Tires"),
    equip("Landscaping Equipment", "Mower — Issues", "Textarea", "No", "", "Free text"),
    equip("Landscaping Equipment", "Mower — Photo", "File Upload", "No", "", ""),
    equip("Landscaping Equipment", "Weed Eater — Gas"),
    equip("Landscaping Equipment", "Weed Eater — Grease"),
    equip("Landscaping Equipment", "Weed Eater — String"),
    equip("Landscaping Equipment", "Weed Eater — Issues", "Textarea", "No", "", ""),
    equip("Landscaping Equipment", "Chainsaw — Gas"),
    equip("Landscaping Equipment", "Chainsaw — Bar Oil"),
    equip("Landscaping Equipment", "Chainsaw — Chain"),
    equip("Landscaping Equipment", "Chainsaw — Issues", "Textarea", "No", "", ""),
    equip("Landscaping Equipment", "Chainsaw — Photo", "File Upload", "No", "", ""),
    equip("Landscaping Equipment", "Blowers — Gas"),
    equip("Landscaping Equipment", "Blowers — Air Filter"),
    equip("Landscaping Equipment", "Blowers — Operational"),
    equip("Landscaping Equipment", "Blowers — Issues", "Textarea", "No", "", ""),
    equip("Landscaping Equipment", "Edgers — Gas"),
    equip("Landscaping Equipment", "Edgers — Blades"),
    equip("Landscaping Equipment", "Edgers — Grease"),
    equip("Landscaping Equipment", "Edgers — Issues", "Textarea", "No", "", ""),
    ("Power Equipment & Vehicles", None, None, None, None, None, None, None, None),
    equip("Power Equipment", "Power Tools — Batteries Charged"),
    equip("Power Equipment", "Power Tools — Blades Good"),
    equip("Power Equipment", "Power Tools — Drill Bits Good"),
    equip("Power Equipment", "Power Tools — Issues", "Textarea", "No", "", ""),
    equip("Power Equipment", "Power Tools — Photo", "File Upload", "No", "", ""),
    equip("Power Equipment", "Golf Carts — Greased"),
    equip("Power Equipment", "Golf Carts — Tires"),
    equip("Power Equipment", "Golf Carts — Gas/Oil"),
    equip("Power Equipment", "Golf Carts — Chargers"),
    equip("Power Equipment", "Golf Carts — Issues", "Textarea", "No", "", ""),
    equip("Power Equipment", "Golf Carts — Photo", "File Upload", "No", "", ""),
    equip("Power Equipment", "Pole Saw — Gas"),
    equip("Power Equipment", "Pole Saw — 2 Chains"),
    equip("Power Equipment", "Pole Saw — Bar Oil"),
    equip("Power Equipment", "Pole Saw — Issues", "Textarea", "No", "", ""),
    ("Pipes & Fittings Inventory", None, None, None, None, None, None, None, None),
    equip("Pipes & Fittings", "PVC Pipe (2 sticks every size)", "Checkbox", "No", "On Hand / Need to Order"),
    equip("Pipes & Fittings", "Glue and Primer (2 each)", "Checkbox", "No", "On Hand / Need to Order"),
    equip("Pipes & Fittings", "Hose Bibs (6 on hand)", "Checkbox", "No", "On Hand / Need to Order"),
    equip("Pipes & Fittings", "Shut-off Valves (6 every size)", "Checkbox", "No", "On Hand / Need to Order"),
    equip("Pipes & Fittings", "Elbows (6 every size)", "Checkbox", "No", "On Hand / Need to Order"),
    equip("Pipes & Fittings", "Slips (6 every size)", "Checkbox", "No", "On Hand / Need to Order"),
    equip("Pipes & Fittings", "Couplings (6 every size)", "Checkbox", "No", "On Hand / Need to Order"),
    equip("Pipes & Fittings", "Plumbing Supply Notes", "Textarea", "No", "", ""),
    equip("Pipes & Fittings", "Photo of Supply Area", "File Upload", "No", "", ""),
    ("Sign-Off", None, None, None, None, None, None, None, None),
    ("Sign-Off", "Overall Equipment Status", "Dropdown", "No", "All Good / Needs Attention / Critical", "None", "Equipment > Overall Status", "Yes", ""),
    ("Sign-Off", "Additional Notes", "Textarea", "No", "", "None", "Equipment > Notes", "Yes", ""),
    ("Sign-Off", "Technician Signature", "Signature", "No", "", "None", "Equipment > Tech Signature", "Yes", ""),
    ("Sign-Off", "Manager Signature", "Signature", "No", "", "None", "Equipment > Mgr Signature", "Yes", ""),
    ("Sign-Off", "Manager Review Date", "Date", "No", "", "None", "Equipment > Review Date", "Yes", ""),
]
add_form_tab(wb, "7-Monthly Equipment", f7)

# FORM 8: Month End Checklist
f8 = [
    ("Community Details", None, None, None, None, None, None, None, None),
    ("Community Details", "Community", "Dropdown", "No", "65-community dropdown", "None", "Month End > Community", "No", ""),
    ("Community Details", "Reporting Month", "Dropdown", "No", "Jan–Dec", "None", "Month End > Month", "No", ""),
    ("Community Details", "Reporting Year", "Dropdown", "No", "Year values", "None", "Month End > Year", "No", ""),
    ("Community Details", "Community Manager Name", "Text", "No", "", "None", "Month End > Manager", "No", ""),
    ("Compliance & Documentation", None, None, None, None, None, None, None, None),
    ("Compliance & Documentation", "Incident Reports Complete", "Dropdown", "No", "Complete / Incomplete", "If Incomplete → explanation field", "Month End > Incidents Complete", "No", ""),
    ("Compliance & Documentation", "Incomplete Explanation", "Textarea", "No", "", "Visible if Incomplete", "Month End > Incidents Notes", "No", ""),
    ("Compliance & Documentation", "Citations/Violations Complete", "Checkbox", "No", "", "None", "Month End > Citations Complete", "No", ""),
    ("Compliance & Documentation", "Violations Updated/Closed", "Checkbox", "No", "", "None", "Month End > Violations Closed", "No", ""),
    ("Compliance & Documentation", "Bill of Sales Uploaded", "Checkbox", "No", "", "None", "Month End > BOS Uploaded", "No", ""),
    ("Compliance & Documentation", "Notes in Rent Manager", "Checkbox", "No", "", "None", "Month End > RM Notes", "No", "OPEN ITEM: RM-only or QB tracked?"),
    ("Compliance & Documentation", "Mileage Submitted", "Checkbox", "No", "", "None", "Month End > Mileage", "No", ""),
    ("Financial & Billing", None, None, None, None, None, None, None, None),
    ("Financial & Billing", "Payments Posted", "Checkbox", "No", "", "None", "Month End > Payments Posted", "No", ""),
    ("Financial & Billing", "Misc Charges Reviewed", "Checkbox", "No", "", "None", "Month End > Misc Charges", "No", ""),
    ("Financial & Billing", "Rent Increases Posted", "Checkbox", "No", "", "None", "Month End > Rent Increases", "No", ""),
    ("Financial & Billing", "Late Fees Posted", "Checkbox", "No", "", "None", "Month End > Late Fees", "No", ""),
    ("Financial & Billing", "Utility Charges Validated", "Checkbox", "No", "", "None", "Month End > Utility Charges", "No", ""),
    ("Resident & Leasing Activity", None, None, None, None, None, None, None, None),
    ("Resident & Leasing", "Notices to Vacate Posted/Uploaded", "Checkbox", "No", "", "None", "Month End > Vacate Notices", "No", ""),
    ("Resident & Leasing", "New Accounts/Move-Ins/Move-Outs Verified", "Checkbox", "No", "", "None", "Month End > Moves Verified", "No", ""),
    ("Resident & Leasing", "Prospects in Rent Manager", "Checkbox", "No", "", "None", "Month End > RM Prospects", "No", "OPEN ITEM: RM-only or QB tracked?"),
    ("Resident & Leasing", "Evictions/Stipulations Posted", "Checkbox", "No", "", "None", "Month End > Evictions", "No", ""),
    ("Resident & Leasing", "Vacant Homes Listed", "Checkbox", "No", "", "None", "Month End > Vacant Listed", "No", ""),
    ("Maintenance & Property", None, None, None, None, None, None, None, None),
    ("Maintenance & Property", "Service Work Orders Closed", "Checkbox", "No", "", "None", "Month End > WOs Closed", "No", ""),
    ("Maintenance & Property", "Maintenance Checklists Signed Off", "Checkbox", "No", "", "None", "Month End > Maint Signed", "No", ""),
    ("Maintenance & Property", "Site Visit Inspections Completed", "Checkbox", "No", "", "None", "Month End > Site Visits", "No", ""),
    ("Maintenance & Property", "Maintenance Tasks Completed/Documented", "Checkbox", "No", "", "None", "Month End > Maint Tasks", "No", ""),
    ("Maintenance & Property", "CapEx Quotes Submitted", "Checkbox", "No", "", "None", "Month End > CapEx Quotes", "No", ""),
    ("Manager Certification", None, None, None, None, None, None, None, None),
    ("Manager Certification", "Community Manager Signature", "Signature", "No", "", "None", "Month End > CM Signature", "No", ""),
    ("Manager Certification", "Date", "Date", "No", "", "None", "Month End > Cert Date", "No", ""),
    ("Manager Certification", "Additional Comments", "Textarea", "No", "", "None", "Month End > Comments", "No", ""),
]
add_form_tab(wb, "8-Month End", f8)

wb.save("/home/ubuntu/parakeet_jotform_audit.xlsx")
print("Part 4 done — Forms 7-8")
